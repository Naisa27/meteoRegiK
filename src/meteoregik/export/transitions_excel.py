"""Выгрузка окна "Переходы температуры" в Excel: две таблицы на двух листах,
как в примере пользователя (.claude/docs/пример_вывода_переходов.xlsx) —
"Таблица переходов" и "Таблица периодов". Раскладка — та же, что на экране
(reports/transitions_view.py); сверху — название станции, как у выгрузки
"Просмотра агрегатных данных".

Значения пишутся текстом, как на экране: в ячейке два прочтения метода
через слэш (Гидрометцентр / "Климат России", корректировка 2026-10-27) —
датой Excel такую ячейку не записать. Порядок прочтений — в строке под
названием станции. "Нет данных" — со штриховкой, как на экране; подсказка
ячейки — в её примечании."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from meteoregik.calculations import transitions as tr
from meteoregik.reports import table_palette
from meteoregik.reports.transitions_view import KIND_NO_DATA, KIND_VALUE, METHODS_ORDER, YEAR_TITLE, Table

SHEET_TRANSITIONS = "Таблица переходов"
SHEET_PERIODS = "Таблица периодов"

# Цвета шапки — те же, что на экране (ui/transitions_dialog.py)
SPRING_FILL = "#DCEFD9"
AUTUMN_FILL = "#F6E3C8"
HEADER_FILL = "#DCE7EE"  # theme.HEADER_BG — экспорт не тянет модуль стилей Qt
HATCH_COLOR = "#B7C2C9"  # светлая: текст "нет данных" поверх штриховки должен читаться

_THIN = Side(style="thin", color=table_palette.THIN_BORDER.lstrip("#"))
_THICK = Side(style="medium", color=table_palette.THICK_BORDER.lstrip("#"))
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HEAD = 3  # ряд групп шапки: выше — название станции и порядок прочтений


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color.lstrip("#"))


def _header_fill(season: str | None) -> PatternFill:
    if season == tr.SPRING:
        return _fill(SPRING_FILL)
    if season == tr.AUTUMN:
        return _fill(AUTUMN_FILL)
    return _fill(HEADER_FILL)


def _write_table(ws: Worksheet, table: Table, station_title: str) -> None:
    width = len(table.columns) + 1
    ws.cell(row=1, column=1, value=station_title).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=1, value=f"В ячейках: {METHODS_ORDER}").font = Font(italic=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    # Шапка в два ряда: группа (весна/осень/длительность) и порог
    group_row, title_row = _HEAD, _HEAD + 1
    year = ws.cell(row=group_row, column=1, value=YEAR_TITLE)
    ws.merge_cells(start_row=group_row, start_column=1, end_row=title_row, end_column=1)
    year.fill = _fill(HEADER_FILL)
    column = 0
    while column < len(table.columns):
        group = table.columns[column].group
        last = column
        while last + 1 < len(table.columns) and table.columns[last + 1].group == group:
            last += 1
        ws.cell(row=group_row, column=column + 2, value=group)
        ws.merge_cells(start_row=group_row, start_column=column + 2, end_row=group_row, end_column=last + 2)
        for index in range(column, last + 1):
            ws.cell(row=group_row, column=index + 2).fill = _header_fill(table.columns[index].season)
        column = last + 1
    for index, col in enumerate(table.columns):
        cell = ws.cell(row=title_row, column=index + 2, value=col.title)
        cell.fill = _header_fill(col.season)
    for row in (group_row, title_row):
        for col in range(1, width + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.alignment = _CENTER
            cell.border = Border(top=_THICK if row == group_row else _THIN,
                                 bottom=_THICK if row == title_row else _THIN, left=_THIN, right=_THIN)

    for offset, row in enumerate(table.rows):
        excel_row = title_row + 1 + offset
        top = _THICK if row.mean else _THIN
        bottom = _THICK if offset == len(table.rows) - 1 else _THIN
        label = ws.cell(row=excel_row, column=1, value=int(row.label) if row.label.isdigit() else row.label)
        label.font = Font(bold=True, color=table_palette.ACCENT_TEXT.lstrip("#"))
        for index, data in enumerate(row.cells):
            cell = ws.cell(row=excel_row, column=index + 2, value=data.text)
            if data.kind == KIND_NO_DATA:
                cell.fill = PatternFill("lightUp", fgColor=HATCH_COLOR.lstrip("#"))
            if row.mean:
                cell.font = Font(bold=True)
            if data.tip and (data.kind != KIND_VALUE or row.mean):
                cell.comment = Comment(data.tip, "meteoRegiK", width=360, height=220)
        for col in range(1, width + 1):
            cell = ws.cell(row=excel_row, column=col)
            cell.alignment = _CENTER
            cell.border = Border(top=top, bottom=bottom, left=_THIN, right=_THIN)

    ws.column_dimensions["A"].width = 16  # "Средняя многолетняя дата" — в две строки
    for index in range(len(table.columns)):
        ws.column_dimensions[get_column_letter(index + 2)].width = 16 if table.columns[index].season else 22
    ws.row_dimensions[title_row].height = 32
    ws.freeze_panes = f"B{title_row + 1}"


def export_transitions(out_path: Path, station_title: str, transitions: Table, periods: Table) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = SHEET_TRANSITIONS
    _write_table(first, transitions, station_title)
    _write_table(workbook.create_sheet(SHEET_PERIODS), periods, station_title)
    workbook.save(out_path)
