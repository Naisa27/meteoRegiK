"""Экспорт результата экрана "Просмотр агрегатных данных" в Excel —
потребляет тот же reports.grid.Grid/reports.chart_data.ChartData, что и
Qt-рендерер экрана (ui/aggregate_table_view.py, ui/aggregate_chart_view.py),
без повторной сборки раскладки.

Экспортируется только текущий видимый на экране режим (одна из "Подробно"/
"Помесячно"/"По годам", либо график) — не все три листа примера сразу
(пример .claude/docs/пример_вывода_агрегатных_данных.xlsx показывает три
раскладки для справки, а не формат единого экспорта). Рядом всегда
откладывается и обратное представление тех же данных: была на экране
таблица — в книгу добавляется график, и наоборот.

Листы (правка 2026-09-25, п.1): график и данные для него — на разных
листах. График — отдельный лист-диаграмма (chartsheet), то есть во весь
лист, с легендой справа от графика и подписанными осями; его ряды ссылаются на
лист "Данные для графика". Порядок: сначала то, что было на экране:
- таблица: "Данные", "График", "Данные для графика";
- график: "График", "Данные для графика".

Осадки (корректировка 2026-10-11) — столбцами, как на экране
(ui/aggregate_chart_view.py); вместе с температурой — на своей оси справа."""

from __future__ import annotations

from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from meteoregik.reports import table_palette
from meteoregik.reports.chart_data import ChartData
from meteoregik.reports.grid import Cell, Fill, Grid
from meteoregik.reports.parameters import ParamKey, decimals_for, is_bar, label_for, quantity_title

_TITLE_FONT = Font(bold=True, size=12)
_HEADER_FONT = Font(bold=True)
# Те же цвета и линии, что и на экране (reports/table_palette.py): тонкая
# сетка только у ячеек таблиц, "medium" — жирный контур.
_THIN_SIDE = Side(style="thin", color=table_palette.THIN_BORDER.lstrip("#"))
_THICK_SIDE = Side(style="medium", color=table_palette.THICK_BORDER.lstrip("#"))
_FILLS: dict[Fill, PatternFill] = {
    "decade": PatternFill("solid", fgColor=table_palette.DECADE_FILL.lstrip("#")),
    "total": PatternFill("solid", fgColor=table_palette.TOTAL_FILL.lstrip("#")),
}


def _number_format(decimals: int) -> str:
    return "0." + "0" * decimals if decimals else "0"


def _write_title(ws: Worksheet, station_title: str, col_count: int) -> None:
    ws.cell(row=1, column=1, value=station_title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(col_count, 1))
    cell = ws.cell(row=1, column=1)
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="center")


def _cell_border(grid_cell: Cell) -> Border | None:
    if not grid_cell.in_table:
        return None

    def side(name: str) -> Side:
        return _THICK_SIDE if name in grid_cell.thick else _THIN_SIDE

    return Border(top=side("top"), bottom=side("bottom"), left=side("left"), right=side("right"))


def _write_grid_sheet(ws: Worksheet, grid: Grid, station_title: str) -> None:
    _write_title(ws, station_title, grid.col_count)

    for row_index, row in enumerate(grid.rows):
        excel_row = row_index + 2
        for col_index, grid_cell in enumerate(row.cells):
            cell = ws.cell(row=excel_row, column=col_index + 1, value=grid_cell.value)
            cell.alignment = Alignment(horizontal=grid_cell.align)
            color = table_palette.ACCENT_TEXT if grid_cell.color == "accent" else table_palette.TEXT
            cell.font = Font(bold=grid_cell.bold, color=color.lstrip("#"))
            if grid_cell.fill is not None:
                cell.fill = _FILLS[grid_cell.fill]
            border = _cell_border(grid_cell)
            if border is not None:
                cell.border = border
            if isinstance(grid_cell.value, float):
                cell.number_format = _number_format(grid_cell.decimals)


def _write_flat_table(ws: Worksheet, chart_data: ChartData, parameters: list[ParamKey], station_title: str) -> int:
    """Плоская таблица "период × показатель" — таблица-источник для
    диаграммы и представление chart_data, когда именно график был основным
    видом на экране. Возвращает номер строки-шапки (нужен вызывающей
    стороне, чтобы построить Reference для нативного графика)."""
    col_count = 1 + len(parameters)
    _write_title(ws, station_title, col_count)

    header_row = 2
    ws.cell(row=header_row, column=1, value=chart_data.x_title).font = _HEADER_FONT
    for i, key in enumerate(parameters):
        cell = ws.cell(row=header_row, column=2 + i, value=label_for(key))
        cell.font = _HEADER_FONT

    for offset, x_label in enumerate(chart_data.x_labels):
        row = header_row + 1 + offset
        ws.cell(row=row, column=1, value=x_label)
        for i, key in enumerate(parameters):
            values = chart_data.series.get(key, [])
            value = values[offset] if offset < len(values) else None
            cell = ws.cell(row=row, column=2 + i, value=value)
            if isinstance(value, float):
                cell.number_format = _number_format(decimals_for(key))

    return header_row


# Размеры шрифтов диаграммы, пт (правка 2026-09-25: значения на осях
# были крупными и слипались).
_TICK_FONT_PT = 8
_LEGEND_FONT_PT = 9
# Подписи оси X — под углом 45° (в 1/60000 градуса, как в OOXML).
_X_LABEL_ROTATION = -45 * 60000
# Не больше стольких подписей по оси X: на "Подробно" за десятилетия (тысячи
# дней) Excel иначе рисует подписи сплошной полосой.
_MAX_X_LABELS = 40
# Область построения (сами кривые, без подписей осей) — доли листа-диаграммы,
# от левого верхнего угла. Слева — место под значения и название оси Y,
# снизу — под наклонные подписи дат и название оси X, справа — под легенду
# (график её не перекрывает). По горизонтали кривые занимают ~80% листа.
_PLOT_AREA = {"x": 0.07, "y": 0.09, "w": 0.80, "h": 0.66}
# С двумя осями значений область уже: между кривыми и легендой нужно место
# под значения и название правой оси (проверено отрисовкой в Excel — при
# ширине 0.80 название оси осадков ложилось на легенду).
_PLOT_AREA_TWO_AXES = {**_PLOT_AREA, "w": 0.71}


def _text_properties(size_pt: int, rotation: int | None = None) -> RichText:
    """Шрифт (и при необходимости поворот) подписей оси/легенды — txPr в
    терминах OOXML; openpyxl не даёт для этого простых свойств."""
    body = RichTextProperties(rot=rotation, vert="horz") if rotation is not None else RichTextProperties()
    char = CharacterProperties(sz=size_pt * 100)
    return RichText(bodyPr=body, p=[Paragraph(pPr=ParagraphProperties(defRPr=char), endParaRPr=char)])


# Идентификатор второй оси значений (у первой — стандартный openpyxl, 100)
_SECOND_AXIS_ID = 200
# Зазор между столбцами, % ширины столбца
_BAR_GAP_WIDTH = 50


def _build_chart(
    data_ws: Worksheet, header_row: int, chart_data: ChartData, parameters: list[ParamKey], station_title: str
) -> LineChart | BarChart:
    """Диаграмма листа "График": кривые и/или столбцы. Возвращается
    диаграмма-контейнер — её заголовок, легенда, раскладка и ось периода
    общие для всех рядов."""
    point_count = len(chart_data.x_labels)
    last_row = header_row + point_count
    # Подписи периода — текст ("01.07.1959"), поэтому ссылка строковая
    # (strRef): set_categories() openpyxl всегда пишет числовую (numRef).
    categories = Reference(data_ws, min_col=1, min_row=header_row + 1, max_row=last_row)

    def filled(chart: LineChart | BarChart, keys: list[ParamKey]) -> LineChart | BarChart:
        # Столбцы листа идут в порядке parameters; ряды одного вида могут
        # стоять не подряд, поэтому каждый добавляется своей ссылкой
        for key in keys:
            column = 2 + parameters.index(key)
            chart.add_data(
                Reference(data_ws, min_col=column, min_row=header_row, max_row=last_row), titles_from_data=True
            )
        for series in chart.series:
            series.cat = AxDataSource(strRef=StrRef(f=str(categories)))
        chart.y_axis.title = quantity_title(keys)
        chart.y_axis.number_format = "0.0"
        # openpyxl по умолчанию пишет у осей delete=True — современный Excel
        # тогда не показывает ни линии осей, ни значения на них.
        chart.y_axis.delete = False
        chart.y_axis.txPr = _text_properties(_TICK_FONT_PT)
        return chart

    line_keys = [key for key in parameters if not is_bar(key)]
    bar_keys = [key for key in parameters if is_bar(key)]
    lines = filled(LineChart(), line_keys) if line_keys else None
    bars = filled(BarChart(), bar_keys) if bar_keys else None
    if bars is not None:
        bars.gapWidth = _BAR_GAP_WIDTH

    if lines is not None and bars is not None:
        # Столбцы — контейнер: ряды рисуются в порядке диаграмм, и кривые
        # должны лечь поверх столбцов. Ось столбцов — справа (пересекает ось
        # периода в её конце), без своей сетки: две сетки не совпали бы.
        bars.y_axis.axId = _SECOND_AXIS_ID
        bars.y_axis.crosses = "max"
        bars.y_axis.majorGridlines = None
        lines.y_axis.axPos = "l"
        plot_area = _PLOT_AREA_TWO_AXES
        chart = bars
        chart += lines  # именно +=: у диаграмм openpyxl это "нарисовать вместе", а + — ошибка типов
    else:
        plot_area = _PLOT_AREA
        chart = lines if lines is not None else bars
        chart.y_axis.axPos = "l"

    chart.title = station_title
    chart.style = 2
    # Легенда справа, вне области построения (overlay=False) — график она не
    # перекрывает; место под неё оставлено в _PLOT_AREA.
    chart.legend.position = "r"
    chart.legend.overlay = False
    chart.legend.txPr = _text_properties(_LEGEND_FONT_PT)
    # chart.layout, не chart.plot_area.layout: при записи openpyxl
    # переписывает plot_area.layout значением chart.layout.
    chart.layout = Layout(
        manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge", **plot_area)
    )

    chart.x_axis.title = chart_data.x_title
    chart.x_axis.delete = False
    # Ось категорий — снизу. openpyxl по умолчанию помечает её как левую
    # (axPos="l", как и ось значений) — Excel тогда путает раскладку, и
    # названия осей налезали на значения.
    chart.x_axis.axPos = "b"
    # Подписи оси X — снизу у края графика, а не у нуля: при отрицательных
    # температурах иначе они налезали бы на кривые.
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.txPr = _text_properties(_TICK_FONT_PT, rotation=_X_LABEL_ROTATION)
    skip = max(1, -(-point_count // _MAX_X_LABELS))  # округление вверх
    chart.x_axis.tickLblSkip = skip
    chart.x_axis.tickMarkSkip = skip
    return chart


def export_aggregate_view(
    *,
    out_path: Path,
    station_title: str,
    primary: Literal["table", "chart"],
    grid: Grid | None,
    chart_data: ChartData,
    parameters: list[ParamKey],
    sheet_titles: tuple[str, str, str] = ("Данные", "График", "Данные для графика"),
) -> None:
    data_sheet_title, chart_sheet_title, chart_data_sheet_title = sheet_titles
    wb = Workbook()
    first_ws = wb.active

    if primary == "table":
        if grid is None:
            raise ValueError("grid обязателен, когда primary == 'table'")
        first_ws.title = data_sheet_title
        _write_grid_sheet(first_ws, grid, station_title)
        chart_data_ws = wb.create_sheet(chart_data_sheet_title)
    else:
        chart_data_ws = first_ws
        chart_data_ws.title = chart_data_sheet_title

    header_row = _write_flat_table(chart_data_ws, chart_data, parameters, station_title)
    # Лист-диаграмма без самой диаграммы Excel считает повреждённым — без
    # данных листа "График" просто нет.
    if chart_data.x_labels and parameters:
        chart_ws = wb.create_chartsheet(chart_sheet_title, 1 if primary == "table" else 0)
        chart_ws.add_chart(_build_chart(chart_data_ws, header_row, chart_data, parameters, station_title))
        if primary == "chart":
            wb.active = 0

    wb.save(out_path)
