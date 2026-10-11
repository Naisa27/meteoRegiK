from __future__ import annotations

import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.chartsheet import Chartsheet

from meteoregik.export.aggregate_excel import export_aggregate_view
from meteoregik.reports import table_palette
from meteoregik.reports.chart_data import ChartData
from meteoregik.reports.grid import Cell, Grid, Row

AVG = ("T", "avg")
MIN = ("T", "min")


def _sample_grid() -> Grid:
    return Grid(
        rows=(
            Row(cells=(Cell("2020 год", align="left", bold=True), Cell())),
            Row(
                cells=(
                    Cell("январь", align="left", bold=True, color="accent", in_table=True, thick=frozenset({"left"})),
                    Cell(6.83, fill="total", in_table=True),
                )
            ),
        ),
        col_count=2,
    )


def _sample_chart_data() -> ChartData:
    return ChartData(x_labels=["2020", "2021"], series={AVG: [5.9, 7.3], MIN: [-10.0, -8.0]}, x_title="Год")


def test_export_table_primary_writes_grid_and_chart_sheets(tmp_path):
    out_path = tmp_path / "export.xlsx"
    export_aggregate_view(
        out_path=out_path,
        station_title="Мосальск",
        primary="table",
        grid=_sample_grid(),
        chart_data=_sample_chart_data(),
        parameters=[AVG, MIN],
    )

    wb = openpyxl.load_workbook(out_path)
    assert wb.sheetnames == ["Данные", "График", "Данные для графика"]

    data_ws = wb["Данные"]
    assert data_ws.cell(row=1, column=1).value == "Мосальск"
    assert data_ws.cell(row=2, column=1).value == "2020 год"
    assert data_ws.cell(row=3, column=2).value == 6.83

    # график — отдельный лист-диаграмма (во весь лист), данные — свой лист
    assert isinstance(wb["График"], Chartsheet)
    chart_data_ws = wb["Данные для графика"]
    assert chart_data_ws.cell(row=1, column=1).value == "Мосальск"
    assert chart_data_ws.cell(row=2, column=1).value == "Год"
    assert chart_data_ws.cell(row=3, column=1).value == "2020"
    assert chart_data_ws._charts == []


def test_export_chart_primary_writes_chart_and_flat_table_sheets(tmp_path):
    out_path = tmp_path / "export.xlsx"
    export_aggregate_view(
        out_path=out_path,
        station_title="Мосальск",
        primary="chart",
        grid=None,
        chart_data=_sample_chart_data(),
        parameters=[AVG, MIN],
    )

    wb = openpyxl.load_workbook(out_path)
    assert wb.sheetnames == ["График", "Данные для графика"]
    assert isinstance(wb["График"], Chartsheet)
    assert wb.active.title == "График"
    assert wb["Данные для графика"].cell(row=3, column=1).value == "2020"


def test_export_table_primary_requires_grid(tmp_path):
    import pytest

    with pytest.raises(ValueError):
        export_aggregate_view(
            out_path=tmp_path / "export.xlsx",
            station_title="Мосальск",
            primary="table",
            grid=None,
            chart_data=_sample_chart_data(),
            parameters=[AVG],
        )


def test_export_table_carries_screen_styles(tmp_path):
    out_path = tmp_path / "export.xlsx"
    export_aggregate_view(
        out_path=out_path,
        station_title="Мосальск",
        primary="table",
        grid=_sample_grid(),
        chart_data=_sample_chart_data(),
        parameters=[AVG, MIN],
    )
    ws = openpyxl.load_workbook(out_path)["Данные"]

    title, month, value = ws.cell(row=2, column=1), ws.cell(row=3, column=1), ws.cell(row=3, column=2)
    # вне таблицы — без границ
    assert title.border.left.style is None and title.border.bottom.style is None
    # жирный контур слева, тонкая сетка по остальным сторонам
    assert month.border.left.style == "medium"
    assert month.border.right.style == "thin"
    assert month.font.bold and month.font.color.rgb.endswith(table_palette.ACCENT_TEXT.lstrip("#"))
    assert value.fill.fgColor.rgb.endswith(table_palette.TOTAL_FILL.lstrip("#"))


def _chart_xml(tmp_path, chart_data: ChartData, parameters) -> ET.Element:
    out_path = tmp_path / "export.xlsx"
    export_aggregate_view(
        out_path=out_path,
        station_title="Мосальск",
        primary="chart",
        grid=None,
        chart_data=chart_data,
        parameters=parameters,
    )
    with zipfile.ZipFile(out_path) as archive:
        return ET.fromstring(archive.read("xl/charts/chart1.xml"))


_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _val(root: ET.Element, path: str) -> str | None:
    node = root.find(path, _NS)
    return None if node is None else node.get("val")


def test_export_chart_legend_right_without_overlapping_plot(tmp_path):
    root = _chart_xml(tmp_path, _sample_chart_data(), [AVG, MIN])
    assert _val(root, ".//c:legend/c:legendPos") == "r"
    assert _val(root, ".//c:legend/c:overlay") == "0"
    layout = root.find(".//c:plotArea/c:layout/c:manualLayout", _NS)
    assert _val(layout, "c:layoutTarget") == "inner"
    x, w = float(_val(layout, "c:x")), float(_val(layout, "c:w"))
    assert w >= 0.75  # график растянут по горизонтали
    assert x + w <= 0.9  # справа остаётся место под легенду


def test_export_chart_axes_positions_fonts_and_titles(tmp_path):
    root = _chart_xml(tmp_path, _sample_chart_data(), [AVG, MIN])
    cat_ax, val_ax = root.find(".//c:catAx", _NS), root.find(".//c:valAx", _NS)
    assert _val(cat_ax, "c:axPos") == "b"
    assert _val(val_ax, "c:axPos") == "l"
    assert _val(cat_ax, "c:delete") == "0" and _val(val_ax, "c:delete") == "0"
    # мелкий шрифт значений осей, подписи оси X под углом
    assert cat_ax.find("c:txPr/a:p/a:pPr/a:defRPr", _NS).get("sz") == "800"
    assert val_ax.find("c:txPr/a:p/a:pPr/a:defRPr", _NS).get("sz") == "800"
    assert cat_ax.find("c:txPr/a:bodyPr", _NS).get("rot") == str(-45 * 60000)
    titles = ["".join(t.text for t in ax.iterfind("c:title//a:t", _NS)) for ax in (cat_ax, val_ax)]
    assert titles == ["Год", "Температура, °C"]


def test_export_chart_categories_are_text_references(tmp_path):
    root = _chart_xml(tmp_path, _sample_chart_data(), [AVG, MIN])
    for series in root.iterfind(".//c:ser", _NS):
        assert series.find("c:cat/c:strRef", _NS) is not None
        assert series.find("c:cat/c:numRef", _NS) is None


def test_export_chart_thins_x_labels_on_long_series(tmp_path):
    labels = [f"{i:05d}" for i in range(4000)]
    chart_data = ChartData(x_labels=labels, series={AVG: [1.0] * 4000}, x_title="Дата")
    root = _chart_xml(tmp_path, chart_data, [AVG])
    assert _val(root, ".//c:catAx/c:tickLblSkip") == "100"  # 4000 / 40 подписей


def test_export_without_chart_points_skips_chart_sheet(tmp_path):
    out_path = tmp_path / "export.xlsx"
    export_aggregate_view(
        out_path=out_path,
        station_title="Мосальск",
        primary="table",
        grid=_sample_grid(),
        chart_data=ChartData(),
        parameters=[AVG],
    )
    assert openpyxl.load_workbook(out_path).sheetnames == ["Данные", "Данные для графика"]


# --- осадки: столбцами, вместе с температурой — на второй оси ---

RAIN = ("R", "sum")


def _chart_data_with_rain() -> ChartData:
    return ChartData(
        x_labels=["2020", "2021"], series={AVG: [5.9, 7.3], RAIN: [612.4, 580.0], MIN: [-10.0, -8.0]}, x_title="Год"
    )


def test_export_chart_draws_precipitation_as_bars_on_second_axis(tmp_path):
    root = _chart_xml(tmp_path, _chart_data_with_rain(), [AVG, RAIN, MIN])

    bar_series = root.findall(".//c:barChart/c:ser", _NS)
    line_series = root.findall(".//c:lineChart/c:ser", _NS)
    assert len(bar_series) == 1 and len(line_series) == 2
    # ряды ссылаются на свои столбцы листа данных: B — t средняя, C — осадки, D — t Мин
    assert [ser.find("c:val/c:numRef/c:f", _NS).text.split("!")[1][1] for ser in line_series] == ["B", "D"]
    assert bar_series[0].find("c:val/c:numRef/c:f", _NS).text.split("!")[1][1] == "C"

    titles = {
        _val(ax, "c:axId"): "".join(t.text for t in ax.iterfind("c:title//a:t", _NS))
        for ax in root.findall(".//c:valAx", _NS)
    }
    assert titles == {"100": "Температура, °C", "200": "Осадки, мм"}
    bar_axis = next(ax for ax in root.findall(".//c:valAx", _NS) if _val(ax, "c:axId") == "200")
    assert _val(bar_axis, "c:crosses") == "max"  # ось осадков — справа
    assert _val(bar_axis, "c:delete") == "0"


def test_export_chart_with_precipitation_only_is_a_bar_chart(tmp_path):
    root = _chart_xml(tmp_path, _chart_data_with_rain(), [RAIN])

    assert len(root.findall(".//c:barChart/c:ser", _NS)) == 1
    assert root.find(".//c:lineChart", _NS) is None
    val_ax = root.find(".//c:valAx", _NS)
    assert "".join(t.text for t in val_ax.iterfind("c:title//a:t", _NS)) == "Осадки, мм"
    assert _val(val_ax, "c:axPos") == "l"


def test_export_formats_precipitation_with_one_decimal(tmp_path):
    out_path = tmp_path / "export.xlsx"
    grid = Grid(rows=(Row(cells=(Cell(5.93, in_table=True), Cell(612.4, in_table=True, decimals=1))),), col_count=2)
    export_aggregate_view(
        out_path=out_path, station_title="Станция", primary="table", grid=grid,
        chart_data=_chart_data_with_rain(), parameters=[AVG, RAIN],
    )

    wb = openpyxl.load_workbook(out_path)
    assert [cell.number_format for cell in wb["Данные"][2]] == ["0.00", "0.0"]
    flat = wb["Данные для графика"]
    assert [flat.cell(row=3, column=col).number_format for col in (2, 3)] == ["0.00", "0.0"]
