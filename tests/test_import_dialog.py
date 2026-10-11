from datetime import date

from openpyxl import Workbook
from sqlalchemy import select
from PySide6.QtWidgets import QApplication, QMessageBox

from meteoregik.calculations.recompute import ValueConflict
from meteoregik.db.migrations import run_migrations
from meteoregik.db.session import get_engine
from meteoregik.importers.base import ImportReport
from meteoregik.db.schema import rp5_readings, stations
from meteoregik.stations.matching import StationNameMismatch
from meteoregik.ui.import_dialog import (
    ImportDialog,
    format_conflict_message,
    format_import_report,
    format_name_mismatch_question,
)


def _qapp():
    return QApplication.instance() or QApplication([])


def _click_button_containing(box: QMessageBox, text: str) -> None:
    for btn in box.buttons():
        if text in btn.text():
            btn.click()
            return
    raise AssertionError(f"кнопка с текстом, содержащим {text!r}, не найдена")


def test_format_import_report_with_period():
    report = ImportReport(
        wmo_index=27704,
        station_name="Мосальск",
        period_start=date(2020, 1, 1),
        period_end=date(2020, 12, 31),
        total_rows=100,
        imported_rows=90,
        skipped_rows=10,
    )

    text = format_import_report(report)

    assert "Мосальск" in text
    assert "2020-01-01 — 2020-12-31" in text
    assert "Всего строк в файле: 100" in text
    assert "Импортировано: 90" in text
    assert "Пропущено (уже существуют): 10" in text


def test_format_import_report_shows_rows_without_data_only_when_known():
    base = dict(
        wmo_index=27704, station_name="Мосальск", period_start=date(2020, 1, 1),
        period_end=date(2020, 1, 2), total_rows=10, imported_rows=6, skipped_rows=1,
    )

    assert "Строк без данных (не загружены): 3" in format_import_report(ImportReport(**base, rows_without_data=3))
    assert "Строк без данных (не загружены): 0" in format_import_report(ImportReport(**base, rows_without_data=0))
    assert "без данных" not in format_import_report(ImportReport(**base))  # rp5: проверки нет


def test_format_import_report_empty_file():
    report = ImportReport(
        wmo_index=None,
        station_name=None,
        period_start=None,
        period_end=None,
        total_rows=0,
        imported_rows=0,
        skipped_rows=0,
    )

    text = format_import_report(report)

    assert "Станция: —" in text
    assert "Период: —" in text


def test_format_conflict_message_day_level():
    conflict = ValueConflict(
        wmo_index=27707,
        period_type="day",
        period_start=date(2015, 1, 21),
        period_end=date(2015, 1, 21),
        stat_type="max",
        old_value=-2.3,
        old_source_file="Сухиничи 1959-2025_формат_ГГ.xlsx",
        new_value=-2.5,
        new_source_file="Сухиничи_2015-2019_rp5.xls",
    )

    text = format_conflict_message(conflict)

    assert "максимальная температура" in text
    assert "2015-01-21" in text
    assert "Сухиничи 1959-2025_формат_ГГ.xlsx" in text
    assert "-2.3" in text
    assert "Сухиничи_2015-2019_rp5.xls" in text
    assert "-2.5" in text


def test_format_conflict_message_without_source_file():
    conflict = ValueConflict(
        wmo_index=27707,
        period_type="decade",
        period_start=date(2015, 1, 1),
        period_end=date(2015, 1, 10),
        stat_type="avg",
        old_value=1.0,
        old_source_file=None,
        new_value=2.0,
        new_source_file=None,
    )

    text = format_conflict_message(conflict)

    assert "неизвестный файл" in text
    assert "2015-01-01 — 2015-01-10" in text


def _make_conflict(day: int, old_value: float, new_value: float) -> ValueConflict:
    return ValueConflict(
        wmo_index=27707,
        period_type="day",
        period_start=date(2015, 1, day),
        period_end=date(2015, 1, day),
        stat_type="max",
        old_value=old_value,
        old_source_file="rp5.xls",
        new_value=new_value,
        new_source_file="gg.xlsx",
    )


def test_conflict_resolver_asks_every_time_without_apply_to_all(tmp_path, monkeypatch):
    _qapp()
    engine = get_engine(tmp_path / "test.sqlite3")
    dialog = ImportDialog(engine)

    call_count = {"n": 0}

    def fake_exec(self):
        call_count["n"] += 1
        _click_button_containing(self, "Взять новое")  # не "для всех"
        return 0

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)

    resolver = dialog._make_conflict_resolver()
    assert resolver(_make_conflict(21, -2.5, -2.3)) == "new"
    assert resolver(_make_conflict(22, 1.0, 2.0)) == "new"
    assert call_count["n"] == 2  # диалог показывался на каждый вызов


def test_conflict_resolver_remembers_apply_to_all_choice(tmp_path, monkeypatch):
    _qapp()
    engine = get_engine(tmp_path / "test.sqlite3")
    dialog = ImportDialog(engine)

    call_count = {"n": 0}

    def fake_exec(self):
        call_count["n"] += 1
        _click_button_containing(self, "Оставить сохранённое для всех оставшихся")
        return 0

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)

    resolver = dialog._make_conflict_resolver()
    assert resolver(_make_conflict(21, -2.5, -2.3)) == "old"
    assert resolver(_make_conflict(22, 1.0, 2.0)) == "old"  # без повторного диалога
    assert resolver(_make_conflict(23, 5.0, 6.0)) == "old"
    assert call_count["n"] == 1  # диалог показан только один раз


def test_run_import_shows_clear_message_for_wrong_format(tmp_path, monkeypatch):
    _qapp()
    engine = get_engine(tmp_path / "test.sqlite3")
    run_migrations(engine)

    # ГГ-файл, но диалог настроен на импорт rp5 — неверный выбор формата.
    wb = Workbook()
    ws = wb.active
    ws.append(["Индекс ВМО", "Год", "Месяц", "День", "Tavg", "Tmax", "RRR", "Tmin"])
    ws.append([27707, 2020, 1, 1, 1.0, 2.0, 0.0, -1.0])
    file_path = str(tmp_path / "gg.xlsx")
    wb.save(file_path)

    dialog = ImportDialog(engine)
    dialog.file_edit.setText(file_path)
    dialog.rp5_radio.setChecked(True)

    calls = []
    monkeypatch.setattr(QMessageBox, "critical", staticmethod(lambda parent, title, text: calls.append((title, text))))

    dialog._run_import()

    assert len(calls) == 1
    title, text = calls[0]
    assert title == "Неверный формат файла"
    assert "неверный формат" in text
    # кнопки снова доступны после ошибки
    assert dialog.import_btn.isEnabled()
    assert dialog.browse_btn.isEnabled()
    assert dialog.close_btn.isEnabled()


def test_run_import_refuses_file_of_station_missing_in_directory(tmp_path, monkeypatch):
    _qapp()
    engine = get_engine(tmp_path / "test.sqlite3")
    run_migrations(engine)  # справочник пуст — станции 27707 в нём нет

    wb = Workbook()
    ws = wb.active
    ws.append(["Индекс ВМО", "Год", "Месяц", "День", "Tavg", "Tmax", "RRR", "Tmin"])
    ws.append([27707, 2020, 1, 1, 1.0, 2.0, 0.0, -1.0])
    file_path = str(tmp_path / "gg.xlsx")
    wb.save(file_path)

    dialog = ImportDialog(engine)
    dialog.file_edit.setText(file_path)
    dialog.gg_radio.setChecked(True)

    calls = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda parent, title, text: calls.append((title, text))))

    dialog._run_import()

    assert len(calls) == 1
    title, text = calls[0]
    assert title == "Станции нет в справочнике"
    assert "27707" in text and "Файл не загружен" in text
    assert "Добавить станцию" in text
    # отчёта об успехе нет, кнопки снова доступны
    assert dialog.success_label.text() == ""
    assert dialog.report_label.text() == ""
    assert dialog.import_btn.isEnabled()


def test_format_name_mismatch_question_names_both_stations_and_the_index():
    text = format_name_mismatch_question(StationNameMismatch(27707, "Мосальск", "Сухиничи"))

    assert "«Мосальск»" in text and "«Сухиничи»" in text and "27707" in text
    assert "записаны станции «Сухиничи»" in text


def _dialog_with_mismatching_rp5_file(tmp_path):
    """В справочнике под индексом 27707 — Сухиничи, а файл подписан Мосальском."""
    engine = get_engine(tmp_path / "test.sqlite3")
    run_migrations(engine)
    with engine.begin() as conn:
        conn.execute(stations.insert().values(name="Сухиничи", wmo_index=27707))
    wb = Workbook()
    ws = wb.active
    ws.append(["# Метеостанция Мосальск, Россия, WMO_ID=27707, выборка с 01.01.2020 по 31.12.2020, все дни"])
    ws.append(["#"])
    ws.append(["Местное время в Мосальске", "T"])
    ws.append(["01.01.2020 09:00", -1.2])
    file_path = str(tmp_path / "rp5.xlsx")
    wb.save(file_path)

    dialog = ImportDialog(engine)
    dialog.file_edit.setText(file_path)
    dialog.rp5_radio.setChecked(True)
    return engine, dialog


def _rp5_row_count(engine) -> int:
    with engine.connect() as conn:
        return len(conn.execute(select(rp5_readings.c.id)).all())


def test_run_import_asks_about_station_name_and_does_not_load_on_refusal(tmp_path, monkeypatch):
    _qapp()
    engine, dialog = _dialog_with_mismatching_rp5_file(tmp_path)
    shown = []

    def fake_exec(self):
        shown.append((self.windowTitle(), self.text(), self.defaultButton().text()))
        _click_button_containing(self, "Не загружать")
        return 0

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)

    dialog._run_import()

    assert len(shown) == 1
    title, text, default_button = shown[0]
    assert title == "Название станции не совпадает"
    assert "«Мосальск»" in text and "«Сухиничи»" in text
    assert default_button == "Не загружать"  # Enter не заливает данные чужой станции
    assert _rp5_row_count(engine) == 0
    assert dialog.status_label.text() == "Файл не загружен"
    assert dialog.success_label.text() == ""
    assert dialog.import_btn.isEnabled()


def test_run_import_loads_file_by_index_after_confirmation(tmp_path, monkeypatch):
    _qapp()
    engine, dialog = _dialog_with_mismatching_rp5_file(tmp_path)

    def fake_exec(self):
        _click_button_containing(self, "Загрузить по индексу")
        return 0

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)

    dialog._run_import()

    assert _rp5_row_count(engine) == 1
    assert dialog.success_label.text() == "Импорт завершён успешно"
    assert "Станция: Сухиничи" in dialog.report_label.text()


def test_import_dialog_buttons_have_rounded_corners():
    _qapp()
    engine = get_engine(":memory:")
    dialog = ImportDialog(engine)

    assert dialog.import_btn.objectName() == "primary"
    assert dialog.browse_btn.objectName() == "secondary"
    assert dialog.close_btn.objectName() == "secondary"
    assert "border-radius" in dialog.styleSheet()


def test_import_dialog_fits_full_report_without_clipping():
    _qapp()
    engine = get_engine(":memory:")
    dialog = ImportDialog(engine)
    dialog.show()
    report = ImportReport(
        27707, "Сухиничи", date(1959, 7, 1), date(2025, 2, 28),
        total_rows=24472, imported_rows=23688, skipped_rows=0, rows_without_data=784,
    )

    dialog.report_label.setText(format_import_report(report))
    QApplication.processEvents()

    assert "Строк без данных (не загружены): 784" in dialog.report_label.text()
    assert dialog.report_label.height() >= dialog.report_label.sizeHint().height()
    assert dialog.report_label.height() >= dialog.report_label.fontMetrics().lineSpacing() * 6

